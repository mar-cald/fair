#!/usr/bin/env python3
"""Read-only FAIR scan of a research project folder.

Usage: python3 fair_scan.py <project_folder> [--out report.json]

Prints a JSON report with: file inventory, naming problems, formats, folder
structure, README coverage, licenses, tabular data checks, data dictionary
coverage, code dependencies and portability problems.

Data files read: CSV/TSV and delimited .txt/.dat (any encoding, UTF-16
included), JSON (jsPsych, JATOS), E-Prime logs (.txt), and Excel or SPSS/Stata
files when openpyxl / pyreadstat are installed. Files from Qualtrics,
PsychoPy, jsPsych, Gorilla, OpenSesame and E-Prime are recognised, and the
rows their exports add above or below the data (question text, ImportId,
"END OF FILE", a file path) are set aside before the checks.

The script never modifies the folder. It prints summaries of the data
(types, ranges, counts), not the rows themselves; columns that look like
personal data are reported without example values.
"""

import argparse
import csv
import io
import json
import os
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from software import EXPERIMENT_EXT, detect_software, experiment_role, is_eprime_log, parse_eprime_log  # noqa: E402

csv.field_size_limit(sys.maxsize)
MAX_BYTES = 50_000_000  # larger text files are scanned on their first 50 MB

SKIP_DIRS = {
    ".git", ".Rproj.user", ".quarto", "node_modules", "__pycache__",
    ".ipynb_checkpoints", ".venv", "venv", "env", ".mypy_cache", ".pytest_cache",
    ".idea", ".vscode", "renv/library", "renv/staging",
}
SKIP_FILES = {".DS_Store", "Thumbs.db", ".Rhistory", ".RData", ".Rapp.history"}

OPEN_TABULAR = {".csv", ".tsv"}
TEXT_DATA = {".txt", ".dat", ".iqdat"}  # read as data when they hold a table, JSON or an E-Prime log
JSON_DATA = {".json", ".jsonl"}
NOT_DATA_JSON = {"dataset_description.json", "package.json", "package-lock.json", "codemeta.json", ".zenodo.json",
                 "tsconfig.json", "manifest.json", "renv.lock", "composer.json", "jatos.json"}
NOT_DATA_TXT = re.compile(r"(?i)^(readme|leggimi|licen[cs]e|copying|notes?|changelog|requirements|citation|install|"
                          r"session_?info|todo|authors|contributors)")
OPEN_TEXT = {".txt", ".json", ".md", ".yaml", ".yml", ".xml", ".cff", ".tex", ".bib", ".html"}
PROPRIETARY_DATA = {
    ".xlsx": "Excel", ".xls": "Excel", ".xlsm": "Excel", ".ods": "OpenDocument spreadsheet (open, but not plain text)",
    ".sav": "SPSS", ".zsav": "SPSS", ".por": "SPSS", ".dta": "Stata", ".sas7bdat": "SAS",
    ".mat": "MATLAB", ".rds": "R binary", ".rdata": "R binary", ".rda": "R binary",
    ".feather": "Arrow/Feather binary", ".parquet": "Parquet binary", ".pkl": "Python pickle",
    ".pickle": "Python pickle", ".h5": "HDF5", ".hdf5": "HDF5 (e.g. PsychoPy ioHub, eye tracking)",
    ".numbers": "Apple Numbers",
    ".edat": "E-Prime data (opens only in E-DataAid)", ".edat2": "E-Prime 2 data (opens only in E-DataAid)",
    ".edat3": "E-Prime 3 data (opens only in E-DataAid)", ".emrg": "E-Prime merged data (E-DataAid)",
    ".emrg2": "E-Prime 2 merged data (E-DataAid)", ".emrg3": "E-Prime 3 merged data (E-DataAid)",
    ".psydat": "PsychoPy pickle (opens only in Python with PsychoPy)",
    ".edf": "EDF binary (EyeLink eye tracking, or EEG)", ".jasp": "JASP analysis file", ".omv": "jamovi analysis file",
}
ARCHIVES = {".zip", ".tar", ".gz", ".tgz", ".7z", ".rar", ".bz2", ".xz"}
PROPRIETARY_DOCS = {".docx": "Word", ".doc": "Word", ".pptx": "PowerPoint", ".ppt": "PowerPoint", ".pages": "Apple Pages", ".odt": "OpenDocument text"}
CODE_EXT = {
    ".r": "R", ".rmd": "R Markdown", ".qmd": "Quarto", ".py": "Python", ".ipynb": "Jupyter",
    ".m": "MATLAB", ".sps": "SPSS syntax", ".do": "Stata", ".jl": "Julia", ".sh": "Shell",
    ".js": "JavaScript", ".sas": "SAS",
}
ENV_FILES = ["renv.lock", "DESCRIPTION", "requirements.txt", "environment.yml", "environment.yaml",
             "pyproject.toml", "Pipfile", "Pipfile.lock", "poetry.lock", "setup.py", "Dockerfile",
             "install.R", "Project.toml", "sessionInfo.txt", "session_info.txt"]

VERSION_WORDS = re.compile(
    r"(?i)(?:^|[ _\-.(])(final[ei]?|def(?:initiv[oaie])?|nuov[oaie]|new|old|vecchi[oaie]?|copy|copia|"
    r"bozza|draft|ultim[oaie]|last|latest|ok|buon[oaie])(?=$|[ _\-.)0-9])"
)
DUPLICATE_MARK = re.compile(r"\(\d+\)|\s\d+$| - Copy| - Copia|_copy\b")
SAFE_NAME = re.compile(r"^[A-Za-z0-9._\-]+$")

MISSING_TOKENS = {"", "na", "n/a", "nan", "null", "none", "undefined", ".", "-", "--", "?", "missing", "mancante",
                  "#n/d", "#n/a", "nd"}
SENTINELS = {"-9", "-99", "-999", "-9999", "99", "999", "9999", "-1", "-77", "-88", "-98"}
BOOL_TOKENS = {"true", "false", "t", "f", "yes", "no", "y", "n", "si", "sì", "vero", "falso"}

PERSONAL_NAME_PATTERNS = [
    (r"(?i)^(first_?name|last_?name|surname|full_?name|nome|cognome|name|nominativo)$", "name"),
    (r"(?i)e_?mail|^mail$|posta", "e-mail"),
    (r"(?i)phone|telefono|cellulare|^tel$|mobile", "phone number"),
    (r"(?i)^ip$|ip_?address|ipaddress", "IP address"),
    (r"(?i)address|indirizzo|^via$|street|^cap$|zip_?code|postcode", "address"),
    (r"(?i)birth|nascita|^dob$|data_?di_?nascita|birthday", "birth date"),
    (r"(?i)codice_?fiscale|^cf$|fiscal|ssn|tax_?code", "tax / national ID"),
    (r"(?i)^(lat|latitude|lon|lng|longitude|location_?latitude|location_?longitude)$", "GPS coordinates"),
    (r"(?i)matricola|student_?id|badge", "institutional ID"),
    (r"(?i)prolific_?pid|^prolific_?id$|mturk|worker_?id|workerid|participant_?public_?id",
     "recruitment-platform ID (e.g. Prolific, MTurk; Gorilla's Public ID)"),
    (r"(?i)^recipient_?(first|last)_?name$", "name"),
    (r"(?i)^external_?(data_?)?reference$", "external reference (an identifier linked to the person)"),
    (r"(?i)^user_?agent$", "browser user-agent string (an indirect identifier)"),
]
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[A-Za-z]{2,}$")
IP_RE = re.compile(r"^(\d{1,3}\.){3}\d{1,3}$")
PROLIFIC_RE = re.compile(r"^[0-9a-f]{24}$")
USER_PATH_RE = re.compile(r"^(/Users/|/home/|[A-Za-z]:[\\/](Users|Documents and Settings)[\\/])")
LIST_RE = re.compile(r"^\[.*\]$", re.S)
OBJ_RE = re.compile(r"^\{.*\}$", re.S)
CF_RE = re.compile(r"^[A-Z]{6}\d{2}[A-Z]\d{2}[A-Z]\d{3}[A-Z]$", re.I)
PHONE_RE = re.compile(r"^\+?[\d\s\-]{9,15}$")
DATE_RE = re.compile(r"^(\d{4}[-/.]\d{1,2}[-/.]\d{1,2}|\d{1,2}[-/.]\d{1,2}[-/.]\d{2,4})([ T]\d{1,2}:\d{2}(:\d{2})?)?$")

ID_LIKE = re.compile(r"(?i)(^id$|_id$|^id_|subj|participant|partecipante|soggetto|^sub$|^pp$|session)")
# columns that identify the participant (for grouping rows), not trials or sessions:
# strong names first, generic "..._id" names only if no strong one is there
PARTICIPANT_ID = re.compile(r"(?i)(^id$|subj|participant|partecipante|soggetto|^sub$|^pp$|^ppn$|response_?id|record_?id|^uid$)")
PARTICIPANT_ID_WEAK = re.compile(r"(?i)(^ids?_|_id$|^code$|^codice$)")
NOT_PARTICIPANT_ID = re.compile(r"(?i)(internal_?node|trial|item|stim|block|experiment|schedule|object|task|condition|node|"
                                r"run|question|page|screen|event|zone|component|study|session|word|image|list)[ _.]?id$|"
                                r"participant[ _.]?(status|starting|completion|device|os|browser|monitor|viewport)")
# columns that describe the participant: one value per person, whatever the number of rows
PERSON_LEVEL = re.compile(r"(?i)^(age|eta|age_years|gender|sex|sesso|genere|education|istruzione|years_of_education|"
                          r"handedness|lateralita|group|gruppo|nationality|nazionalita|native_language|language|lingua|occupation)$")
SYNONYMS = {"eta": "age", "age_years": "age", "sesso": "sex", "genere": "gender", "gruppo": "group",
            "istruzione": "education", "lateralita": "handedness", "nazionalita": "nationality", "lingua": "language"}
DICT_NAME = re.compile(r"(?i)(dictionary|codebook|dizionario|data_?dict|variables|variabili|metadata)")
DICT_COLUMNS = ["variable_name", "label", "type", "unit", "allowed_values", "coding", "missing_values", "source", "notes"]
DICT_ALIASES = {
    "variable_name": ["variable_name", "variable", "var", "name", "column", "colonna", "variabile", "nome_variabile", "field"],
    "label": ["label", "description", "descrizione", "etichetta", "meaning", "significato"],
    "type": ["type", "tipo", "data_type", "datatype", "class"],
    "unit": ["unit", "units", "unita", "unità", "unit_of_measure"],
    "allowed_values": ["allowed_values", "range", "values", "valori", "valid_values", "levels"],
    "coding": ["coding", "codes", "codifica", "value_labels", "levels_labels"],
    "missing_values": ["missing_values", "missing", "na", "mancanti", "valori_mancanti"],
    "source": ["source", "fonte", "origin", "instrument", "strumento"],
    "notes": ["notes", "note", "comments", "commenti"],
}

# functions that draw random numbers, and ways of fixing the seed
RANDOM_R = re.compile(
    r"\b(sample|sample_n|sample_frac|slice_sample|rnorm|runif|rbinom|rpois|rgamma|rbeta|rexp|rchisq|rmultinom|rlnorm|"
    r"rweibull|rnbinom|rmvnorm|mvrnorm|brm|stan|sampling|stan_glm|stan_glmer|stan_lmer|stan_polr|ulam|map2stan|"
    r"MCMCglmm|jags|jags\.model|run\.jags|coda\.samples|boot|bootMer|bootstrap|bootstrap_model|bootstrap_parameters|"
    r"mice|kmeans|anovaBF|ttestBF|lmBF|generalTestBF|regressionBF|correlationBF|contingencyTableBF|proportionBF|"
    r"posterior|bridge_sampler|posterior_predict|posterior_epred|posterior_linpred|pp_check|blavaan|bsem|bcfa|bgrowth|"
    r"simulate|permute|shuffle)\s*\(|\$sample\s*\(")
SEED_R = re.compile(r"set\.seed\s*\(\s*[^)\s][^)]*\)|with_seed\s*\(|local_seed\s*\(|\bseed\s*=\s*[\w.]+")
RANDOM_PY = re.compile(
    r"np\.random\.\w+|numpy\.random\.\w+|\brandom\.(random|randint|choice|choices|shuffle|sample|uniform|gauss|normalvariate)\s*\(|"
    r"default_rng\s*\(|\b(pm|pymc|pymc3)\.sample\s*\(|\bbootstrap\s*\(|\bresample\s*\(|train_test_split\s*\(|"
    r"KMeans\s*\(|torch\.rand\w*\s*\(|\.sample\s*\(")
SEED_PY = re.compile(r"\.seed\s*\(\s*\d|default_rng\s*\(\s*\d|RandomState\s*\(\s*\d|manual_seed\s*\(|"
                     r"random_state\s*=\s*\d|random_seed\s*=\s*\d|\bseed\s*=\s*\d")
RANDOM_M = re.compile(r"\b(rand|randn|randi|randperm|bootstrp|datasample)\s*\(")
SEED_M = re.compile(r"\brng\s*\(")

R_BASE = {"base", "stats", "utils", "graphics", "grDevices", "methods", "datasets", "tools",
          "parallel", "grid", "splines", "stats4", "tcltk", "compiler"}

README_SECTIONS = {
    "title": None,  # checked separately: first heading / first line
    "authors_orcid": r"(?i)(author|autor|autrice|orcid|contact|contatt|corresponding|maintainer)",
    "description": r"(?i)(about|description|descrizione|overview|study|studio|abstract|panoramica|di cosa)",
    "contents": r"(?i)(contents?|contenut|files?\b|structure|struttura|organi[sz]|folder|cartell)",
    "how_to_read": r"(?i)(how to read|reading the data|come leggere|data format|formato dei dati|codebook|dictionary|dizionario)",
    "how_to_reproduce": r"(?i)(reproduc|riprodu|how to run|usage|utilizzo|requirements|requisiti|install|software|dependencies|dipendenze|renv|sessioninfo)",
    "license": r"(?i)(licen[sc]e|licenza)",
    "citation": r"(?i)(cit(e|ation|are|azione)|doi\.org|10\.\d{4,9}/)",
    "related_materials": r"(?i)(related|links?\b|collegat|preregistr|article|articolo|paper|publication|pubblicazion|preprint)",
    "provenance": r"(?i)(collected|raccolt|recruited|reclutat|participants|partecipanti|sample|campione|data collection|between .* and|fra .* e|tra .* e)",
}

LICENSE_SIGNATURES = [
    ("CC0-1.0", [r"CC0 1\.0", r"Creative Commons Legal Code\s+CC0", r"publicdomain/zero", r"\bCC0\b"]),
    ("CC-BY-NC-ND-4.0", [r"Attribution-NonCommercial-NoDerivatives", r"by-nc-nd/4", r"CC[ -]BY[ -]NC[ -]ND"]),
    ("CC-BY-NC-SA-4.0", [r"Attribution-NonCommercial-ShareAlike", r"by-nc-sa/4", r"CC[ -]BY[ -]NC[ -]SA"]),
    ("CC-BY-NC-4.0", [r"Attribution-NonCommercial 4\.0", r"licenses/by-nc/4", r"CC[ -]BY[ -]NC(?![ -](SA|ND))"]),
    ("CC-BY-ND-4.0", [r"Attribution-NoDerivatives 4\.0", r"licenses/by-nd/4", r"CC[ -]BY[ -]ND"]),
    ("CC-BY-SA-4.0", [r"Attribution-ShareAlike 4\.0", r"licenses/by-sa/4", r"CC[ -]BY[ -]SA"]),
    ("CC-BY-4.0", [r"Attribution 4\.0 International", r"licenses/by/4\.0", r"CC[ -]BY[ -]?4(\.0)?(?![ -]?(NC|SA|ND))", r"CC[ -]BY(?![ -]?(NC|SA|ND|-?\d))"]),
    ("AGPL-3.0", [r"GNU AFFERO GENERAL PUBLIC LICENSE", r"\bAGPL"]),
    ("LGPL-3.0", [r"GNU LESSER GENERAL PUBLIC LICENSE", r"\bLGPL"]),
    ("GPL-3.0", [r"GNU GENERAL PUBLIC LICENSE\s+Version 3", r"\bGPL[ -]?(v)?3", r"GPL-3\.0"]),
    ("GPL-2.0", [r"GNU GENERAL PUBLIC LICENSE\s+Version 2", r"\bGPL[ -]?(v)?2"]),
    ("MIT", [r"\bMIT License\b", r"Permission is hereby granted, free of charge", r"\bMIT\b"]),
    ("Apache-2.0", [r"Apache License,?\s+Version 2\.0", r"Apache[ -]2\.0"]),
    ("BSD", [r"Redistribution and use in source and binary forms", r"\bBSD[ -]?[023]"]),
    ("0BSD", [r"\b0BSD\b", r"Zero-Clause BSD"]),
    ("scientific-use-only", [r"(?i)scientific use file", r"(?i)scientific purposes only", r"(?i)uso scientifico"]),
]


def rel(p, root):
    return str(Path(p).relative_to(root))


def walk(root):
    files, dirs = [], []
    for dirpath, dirnames, filenames in os.walk(root):
        relp = os.path.relpath(dirpath, root)
        keep = []
        for d in sorted(dirnames):
            r = os.path.normpath(os.path.join(relp, d)) if relp != "." else d
            if d in SKIP_DIRS or r in SKIP_DIRS or d.startswith(".") and d not in {".github"}:
                continue
            keep.append(d)
            dirs.append(r)
        dirnames[:] = keep
        for f in sorted(filenames):
            if f in SKIP_FILES or f.startswith("~$") or f.startswith("._"):
                continue
            files.append(Path(dirpath) / f)
    return files, dirs


def name_problems(name, is_dir=False):
    probs = []
    stem = name if is_dir else Path(name).stem
    if " " in name:
        probs.append("spaces")
    if any(unicodedata.category(c).startswith("L") and ord(c) > 127 for c in name):
        probs.append("accented/non-ASCII letters")
    elif not SAFE_NAME.match(name.replace(" ", "_")):
        bad = sorted({c for c in name if not re.match(r"[A-Za-z0-9._\- ]", c)})
        probs.append("special characters " + "".join(bad))
    if VERSION_WORDS.search(stem):
        probs.append("manual versioning word (" + VERSION_WORDS.search(stem).group(1) + ")")
    if DUPLICATE_MARK.search(stem):
        probs.append("looks like a duplicate copy")
    return probs


def check_numbering(files, root):
    """Numbered files in the same folder without leading zeros."""
    by_dir = defaultdict(list)
    for f in files:
        m = re.match(r"^(\d+)[_\-. ]", f.name)
        if m:
            by_dir[f.parent].append((m.group(1), f.name))
    out = []
    for d, items in by_dir.items():
        widths = {len(n) for n, _ in items}
        if len(items) >= 2 and (len(widths) > 1 or (widths == {1} and len(items) >= 2)):
            out.append({"folder": rel(d, root) if d != root else ".",
                        "files": [n for _, n in sorted(items)],
                        "problem": "numbers without leading zeros (use 01_, 02_ ...)"})
    return out


# ---------------------------------------------------------------- tabular data

def decode_bytes(raw):
    """Bytes of a text file -> (text, encoding, has_bom). Handles UTF-16 (E-Prime
    logs and E-DataAid exports, Qualtrics TSV), with or without a BOM."""
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw[3:].decode("utf-8", errors="replace"), "utf-8", True
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16", errors="replace").lstrip("﻿"), "utf-16", True
    head = raw[:4000]
    if head and head.count(b"\x00") > len(head) // 4:
        enc = "utf-16-le" if head[1::2].count(0) > head[0::2].count(0) else "utf-16-be"
        return raw.decode(enc, errors="replace"), enc, False
    for enc in ("utf-8", "cp1252"):
        try:
            return raw.decode(enc), enc, False
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1"), "latin-1", False


def read_text(path):
    """Text of a file and how it is encoded. Files above MAX_BYTES are read up to
    the last full line before that size, and say so in info["sampled"]."""
    size = path.stat().st_size
    with open(path, "rb") as fh:
        raw = fh.read(MAX_BYTES if size > MAX_BYTES else -1)
    info = {}
    if size > MAX_BYTES:
        utf16 = raw[:2] in (b"\xff\xfe", b"\xfe\xff") or raw[:4000].count(b"\x00") > 1000
        nl = b"\n\x00" if utf16 and raw[:2] != b"\xfe\xff" else (b"\x00\n" if utf16 else b"\n")
        cut = raw.rfind(nl)
        if cut > 0:
            raw = raw[:cut + len(nl)]
        info["sampled"] = f"only the first {MAX_BYTES // 1_000_000} MB of {size / 1e6:.0f} MB were read"
    text, enc, bom = decode_bytes(raw)
    info.update({"encoding": enc, "bom": bom})
    if "\r\n" in text[:5000]:
        info["line_endings"] = "CRLF"
    return text, info


DELIMS = {",": "comma", ";": "semicolon", "\t": "tab", "|": "pipe"}


def parse_delimited(text, suffix=".csv"):
    """Parse delimited text. Returns (rows, info) or (None, info) when the text is
    not a table. Quoted cells may span lines (free-text answers). Lines above the
    header with a single field (E-DataAid writes the .edat path there) are set
    aside in info["preamble"]."""
    sample_lines = text[:50000].splitlines()[:40]
    sample = "\n".join(sample_lines)
    best = None
    for d in ([("\t")] if suffix == ".tsv" else list(DELIMS)):
        try:
            counts = [len(r) for r in csv.reader(io.StringIO(sample), delimiter=d) if any(c.strip() for c in r)]
        except csv.Error:
            continue
        if not counts:
            continue
        mode, n_mode = Counter(counts).most_common(1)[0]
        if mode < 2:
            continue
        share = n_mode / len(counts)
        key = (round(share, 1), mode)
        if best is None or key > best[0]:
            best = (key, d, mode, share)
    info = {}
    if best is None:
        return None, info
    _, delim, mode, share = best
    info["delimiter"] = DELIMS[delim]
    rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    pre = []
    while rows and len(pre) < 5 and len(rows) > 2 and len([c for c in rows[0] if c.strip()]) <= 1 and len(rows[1]) >= max(2, mode // 2):
        pre.append(" ".join(rows.pop(0)).strip()[:160])
    if pre:
        info["preamble"] = pre
    info["_share"] = share
    return rows, info


def read_table(path):
    """Read a CSV/TSV or delimited text file: (rows, info). Raises ValueError if
    the file does not hold a table."""
    text, info = read_text(path)
    rows, extra = parse_delimited(text, path.suffix.lower())
    if rows is None:
        if not text.strip():
            return [], {**info, "delimiter": "comma"}
        # a single column: still a table
        rows = list(csv.reader(io.StringIO(text)))
        extra = {"delimiter": "comma"}
    extra.pop("_share", None)
    info.update(extra)
    return rows, info


def _cell(v):
    """A JSON / Excel / SPSS value as the text it would have in a CSV."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, float):
        if v != v:
            return ""
        return str(int(v)) if v.is_integer() and abs(v) < 1e15 else repr(v)
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return str(v)


def json_table(text):
    """Records in a JSON data file -> rows, or None if it is not a table of records.

    Accepts an array of objects (jsPsych's .json()), an object holding such an
    array ("trials", "data", "values", "results"), and JSON lines with one object
    or array per line (JATOS result files)."""
    text = text.strip()
    try:
        items = [json.loads(text)]
    except json.JSONDecodeError as whole:
        items = []
        for line in text.splitlines():
            if not line.strip():
                continue
            try:
                items.append(json.loads(line))
            except json.JSONDecodeError:
                raise ValueError(f"not valid JSON ({whole.msg}, line {whole.lineno}): truncated or damaged?") from None
    records = []

    def collect(o, depth=0):
        if depth > 4:
            return
        if isinstance(o, list):
            for x in o:
                if isinstance(x, dict):
                    records.append(x)
                elif isinstance(x, list):
                    collect(x, depth + 1)
        elif isinstance(o, dict):
            for key in ("trials", "data", "values", "results", "responses"):
                v = o.get(key)
                if isinstance(v, list) and v and all(isinstance(x, dict) for x in v[:20]):
                    collect(v, depth + 1)
                    return
            records.append(o)

    for it in items:
        collect(it)
    if len(records) < 2:
        return None
    cols = []
    for r in records:
        for k in r:
            if k not in cols:
                cols.append(k)
    if len(cols) < 2:
        return None
    return [cols] + [[_cell(r.get(k)) for k in cols] for r in records]


def load_xlsx(path, max_rows=200_000):
    """Sheets of an Excel file as [(sheet, rows)], or raises ImportError/ValueError."""
    import openpyxl  # noqa: PLC0415
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = []
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                break
            rows.append([_cell(v) for v in r])
        rows = [r for r in rows if any(c.strip() for c in r)]
        if not rows:
            continue
        width = max((max((j for j, c in enumerate(r) if c.strip()), default=-1) for r in rows), default=-1) + 1
        out.append((ws.title, [r[:width] + [""] * (width - len(r[:width])) for r in rows]))
    wb.close()
    return out


def load_stat(path):
    """SPSS / Stata / SAS file -> (rows, labels) with pyreadstat, or raises ImportError."""
    import pyreadstat  # noqa: PLC0415
    suf = path.suffix.lower()
    reader = {".sav": pyreadstat.read_sav, ".zsav": pyreadstat.read_sav, ".por": pyreadstat.read_por,
              ".dta": pyreadstat.read_dta, ".sas7bdat": pyreadstat.read_sas7bdat}[suf]
    df, meta = reader(str(path))
    rows = [list(map(str, df.columns))] + [[_cell(v) for v in rec] for rec in df.itertuples(index=False, name=None)]
    labels = {"variable_labels": {k: v for k, v in zip(meta.column_names, meta.column_labels or []) if v},
              "value_labels": {k: {_cell(a): b for a, b in v.items()} for k, v in (meta.variable_value_labels or {}).items()}}
    return rows, labels


def load_any(path):
    """Every table a data file holds: a list of (sheet or None, rows, info).
    Empty list: not a data file. rows None: a data file that could not be read
    here, with the reason in info["not_read"]."""
    suf, name = path.suffix.lower(), path.name
    if suf in OPEN_TABULAR:
        rows, info = read_table(path)
        return [(None, rows, info)]
    if suf in JSON_DATA or suf in TEXT_DATA:
        if name.lower() in NOT_DATA_JSON or (suf in TEXT_DATA and NOT_DATA_TXT.match(name)):
            return []
        text, info = read_text(path)
        if suf in TEXT_DATA and is_eprime_log(text):
            table, ep = parse_eprime_log(text)
            info.update({"format": "E-Prime log (.txt written during the session)",
                         "eprime_levels": ep["levels"], "rows_are": ep["row_level"]})
            return [(None, table, info)]
        if text.lstrip()[:1] in ("[", "{"):
            if info.get("sampled"):
                return [(None, None, {**info, "not_read": "JSON file too large to scan"})]
            try:
                table = json_table(text)
            except ValueError as e:
                return [(None, None, {**info, "not_read": str(e)})]
            if table:
                info["format"] = "JSON records" if suf != ".txt" else "JSON records in a .txt (e.g. a JATOS result file)"
                return [(None, table, info)]
            return []
        if suf in JSON_DATA:
            return []
        rows, extra = parse_delimited(text, suf)
        if rows is None or extra.get("_share", 0) < 0.8 or len(rows) < 3:
            return []
        extra.pop("_share", None)
        info.update(extra)
        return [(None, rows, info)]
    if suf in (".xlsx", ".xlsm"):
        try:
            return [(sheet, rows, {"format": "Excel", "read_with": "openpyxl"}) for sheet, rows in load_xlsx(path)]
        except ImportError:
            return [(None, None, {"not_read": "openpyxl not installed: content not checked"})]
        except Exception as e:  # noqa: BLE001
            return [(None, None, {"not_read": f"could not read: {e}"})]
    if suf in (".sav", ".zsav", ".por", ".dta", ".sas7bdat"):
        try:
            rows, labels = load_stat(path)
            return [(None, rows, {"format": PROPRIETARY_DATA[suf], "read_with": "pyreadstat",
                                  "n_variable_labels": len(labels["variable_labels"]),
                                  "n_value_labels": len(labels["value_labels"])})]
        except ImportError:
            return [(None, None, {"not_read": "pyreadstat not installed: content not checked (R: haven::read_sav)"})]
        except Exception as e:  # noqa: BLE001
            return [(None, None, {"not_read": f"could not read: {e}"})]
    return []


def load_table(path, sheet=None):
    """The one table in a data file (or the sheet asked for): (rows, info).
    Used by the other scripts; raises ValueError when there is no table."""
    path = Path(path)
    found = load_any(path)
    if sheet is not None:
        found = [f for f in found if f[0] == sheet]
    found = [f for f in found if f[1] is not None]
    if not found:
        if path.suffix.lower() in OPEN_TABULAR:
            return read_table(path)
        raise ValueError(f"{path.name}: no table found" + (f" in sheet {sheet}" if sheet else ""))
    return found[0][1], found[0][2]


END_MARK = re.compile(r"(?i)^\s*(end of file|eof)\s*$")


def split_export(rows):
    """Separate what an export adds around the data from the data itself.

    Returns (header, body, labels, layout): `labels` are the rows under the
    header that describe the columns (Qualtrics: question text and ImportId;
    Gorilla's optional question-text row), `layout` says what was set aside:
    label rows, footer rows ("END OF FILE"), an empty unnamed last column."""
    rows = [r for r in rows if any(c.strip() for c in r)]
    layout = {}
    if not rows:
        return [], [], [], layout
    header, body = rows[0], rows[1:]
    # footer
    n_footer = 0
    while body and n_footer < 3 and any(END_MARK.match(c) for c in body[-1] if c.strip()) and \
            len([c for c in body[-1] if c.strip()]) <= 2:
        body.pop()
        n_footer += 1
    if n_footer:
        layout["footer_rows"] = n_footer
    # label rows right under the header
    n_labels = 0
    if sum(bool(re.match(r"^V\d+$", h.strip())) for h in header) >= 3 and body and \
            any(c.strip().lower() == "responseid" for c in body[0]):
        n_labels = 1  # Qualtrics legacy export
    for i in range(min(3, len(body))):
        if any(c.strip().startswith('{"ImportId"') for c in body[i]):
            n_labels = i + 1
    if not n_labels:
        for i in range(min(2, max(0, len(body) - 3))):
            rest = body[i + 1:i + 301]
            numeric = []
            for j in range(len(header)):
                vals = [r[j] for r in rest if j < len(r) and r[j].strip() and r[j].strip().lower() not in MISSING_TOKENS]
                if len(vals) >= 3 and sum(to_number(v) is not None for v in vals) >= 0.9 * len(vals):
                    numeric.append(j)
            row = body[i]
            text_in_numeric = [j for j in numeric if j < len(row) and row[j].strip()
                               and row[j].strip().lower() not in MISSING_TOKENS and to_number(row[j]) is None]
            if len(numeric) >= 2 and len(text_in_numeric) >= 0.6 * len(numeric):
                n_labels = i + 1
            else:
                break
    labels = body[:n_labels]
    body = body[n_labels:]
    if n_labels:
        layout["label_rows"] = n_labels
        if any('{"ImportId"' in c for r in labels for c in r):
            layout["label_rows_hold"] = "question text and ImportId (Qualtrics)"
        elif any(c.strip().lower() == "responseid" for c in labels[0]):
            layout["label_rows_hold"] = "variable names and question text (Qualtrics legacy export)"
        else:
            layout["label_rows_hold"] = "labels such as the question text"
    # an unnamed, empty last column (PsychoPy ends every row with a comma)
    if len(header) > 1 and not header[-1].strip() and all(len(r) < len(header) or not r[len(header) - 1].strip() for r in body):
        header = header[:-1]
        body = [r[:len(header)] for r in body]
        labels = [r[:len(header)] for r in labels]
        layout["empty_last_column"] = True
    return header, body, labels, layout


def to_number(v, decimal_comma=False):
    s = v.strip()
    if decimal_comma:
        s = s.replace(".", "").replace(",", ".") if re.match(r"^-?\d{1,3}(\.\d{3})*(,\d+)?$", s) else s.replace(",", ".")
    try:
        x = float(s)
        if x != x:  # nan
            return None
        return x
    except ValueError:
        return None


def infer_column(name, values, decimal_comma):
    n = len(values)
    stripped = [v.strip() for v in values]
    missing = Counter(v for v in stripped if v.lower() in MISSING_TOKENS)
    present = [v for v in stripped if v.lower() not in MISSING_TOKENS]
    col = {"name": name, "n_missing": sum(missing.values()),
           "missing_tokens": {(k if k else "<empty>"): c for k, c in missing.items()}}
    issues = []
    if len(missing) > 1:
        issues.append("missing values coded in more than one way: " + ", ".join(col["missing_tokens"]))
    ws = sum(1 for v in values if v != v.strip() and v.strip())
    if ws:
        issues.append(f"{ws} values with leading/trailing spaces")
    if not present:
        col["type"] = "empty"
        issues.append("column is entirely empty")
        col["issues"] = issues
        return col

    nums = [to_number(v, decimal_comma) for v in present]
    num_ok = [x for x in nums if x is not None]
    lower = [v.lower() for v in present]
    uniq = Counter(present)
    col["n_unique"] = len(uniq)
    n_list = sum(1 for v in present if LIST_RE.match(v))
    n_obj = sum(1 for v in present if OBJ_RE.match(v))

    if n_list + n_obj >= 0.5 * len(present):
        col["type"] = "JSON object" if n_obj >= n_list else "list"
        keys = Counter()
        for v in present[:2000]:
            try:
                o = json.loads(v)
            except (json.JSONDecodeError, ValueError):
                continue
            if isinstance(o, dict):
                keys.update(o.keys())
        if keys:
            col["json_keys"] = [k for k, _ in keys.most_common(30)]
        issues.append("cells hold " + ("JSON objects" if col["type"] == "JSON object" else "lists") +
                      " (several values per cell): describe their structure in the dictionary, or unpack them into columns")
    elif n_list + n_obj >= max(2, 0.1 * len(present)):
        keys = Counter()
        for v in present:
            if OBJ_RE.match(v):
                try:
                    o = json.loads(v)
                except (json.JSONDecodeError, ValueError):
                    continue
                if isinstance(o, dict):
                    keys.update(o.keys())
        if keys:
            col["json_keys"] = [k for k, _ in keys.most_common(30)]
        issues.append(f"{n_list + n_obj} of {len(present)} cells hold lists or JSON objects (e.g. survey answers) "
                      "and the rest single values: describe both in the dictionary")
    if col.get("type"):
        pass
    elif all(l in BOOL_TOKENS for l in lower) and len(set(lower)) <= 2:
        col["type"] = "boolean"
        col["values"] = sorted(uniq)
    elif len(num_ok) == len(present):
        col["type"] = "integer" if all(float(x).is_integer() for x in num_ok) else "numeric"
        col["min"], col["max"] = min(num_ok), max(num_ok)
        if col["type"] == "integer":
            col["min"], col["max"] = int(col["min"]), int(col["max"])
        sent = {v: c for v, c in uniq.items() if v in SENTINELS}
        if sent and len(uniq) > 3:
            vals_wo = [x for x, v in zip(nums, present) if v not in SENTINELS and x is not None]
            if vals_wo and any(abs(float(s)) > 3 * max(abs(min(vals_wo)), abs(max(vals_wo)), 1) or float(s) < 0 <= min(vals_wo) for s in sent):
                issues.append("possible numeric missing-value codes: " + ", ".join(f"{k} (x{c})" for k, c in sent.items()))
        if col["type"] == "integer" and len(uniq) <= 10:
            col["values"] = sorted(uniq, key=lambda s: float(to_number(s, decimal_comma)))
        if not decimal_comma and any("," in v for v in present):
            issues.append("decimal comma inside a comma-delimited file")
    elif len(num_ok) >= 0.8 * len(present):
        col["type"] = "mixed"
        bad = Counter(v for v, x in zip(present, nums) if x is None)
        issues.append("mostly numeric but contains text: " + ", ".join(f"'{k}' (x{c})" for k, c in bad.most_common(5)))
    elif all(DATE_RE.match(v) for v in present):
        col["type"] = "date"
        formats = {re.sub(r"\d", "9", v) for v in present}
        if len(formats) > 1:
            issues.append("dates written in more than one format: " + ", ".join(sorted(formats)[:4]))
    else:
        col["type"] = "categorical" if len(uniq) <= max(12, 0.05 * n) else "string"
        if col["type"] == "categorical":
            col["values"] = sorted(uniq)
            folded = defaultdict(set)
            for v in uniq:
                folded[v.lower().strip()].add(v)
            clash = [sorted(s) for s in folded.values() if len(s) > 1]
            if clash:
                issues.append("same category written differently: " + "; ".join(" / ".join(c) for c in clash))
    if len(uniq) == 1 and n > 1 and not ID_LIKE.search(name):
        issues.append("constant column (one value only)")
    col["issues"] = issues
    return col


def personal_flags(name, values, label=""):
    """Why a column may hold personal data, from its name, its label (e.g. the
    question text of a Qualtrics export) and its values."""
    flags = []
    for nm in (name, label):
        nm = re.sub(r"[ .\-]+", "_", (nm or "").strip())
        if not nm or len(nm) > 40:
            continue
        hit = next((lab for pat, lab in PERSONAL_NAME_PATTERNS if re.search(pat, nm)), None)
        if hit:
            flags.append(f"column name suggests {hit}" if nm == re.sub(r"[ .\-]+", "_", name.strip())
                         else f"its label ('{label.strip()[:40]}') suggests {hit}")
            break
    vals = [v.strip() for v in values if v.strip()][:2000]
    if vals:
        if sum(bool(EMAIL_RE.match(v)) for v in vals) > 0.3 * len(vals):
            flags.append("values look like e-mail addresses")
        if sum(bool(CF_RE.match(v)) for v in vals) > 0.3 * len(vals):
            flags.append("values look like Italian tax codes (codice fiscale)")
        if sum(bool(PHONE_RE.match(v)) and len(re.sub(r"\D", "", v)) >= 9 for v in vals) > 0.5 * len(vals) and not re.search(r"(?i)(rt|time|ms|id|code)", name):
            flags.append("values look like phone numbers")
        if sum(bool(IP_RE.match(v)) for v in vals) > 0.3 * len(vals):
            flags.append("values look like IP addresses")
        if sum(bool(PROLIFIC_RE.match(v)) for v in vals) > 0.5 * len(vals):
            flags.append("values look like Prolific participant IDs")
        if sum(bool(USER_PATH_RE.match(v)) for v in vals) > 0.3 * len(vals):
            flags.append("values are file paths of a computer, with a user name in them")
        uniq = set(vals)
        texts = []
        for v in uniq:
            if OBJ_RE.match(v) or LIST_RE.match(v):
                try:
                    o = json.loads(v)
                except (json.JSONDecodeError, ValueError):
                    continue
                leaves = o.values() if isinstance(o, dict) else o if isinstance(o, list) else []
                texts += [x for x in leaves if isinstance(x, str)]
            else:
                texts.append(v)
        long_text = [v for v in set(texts) if len(v) > 40 and " " in v]
        if len(long_text) >= 3 and len(long_text) >= 0.3 * len(set(texts) or {1}):
            flags.append("free-text answers: read them before sharing, they can contain names or other identifying details")
    return flags


def canon_name(name):
    folded = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode()
    n = re.sub(r"[^a-z0-9]+", "_", folded.strip().lower()).strip("_")
    return SYNONYMS.get(n, n)


def is_missing(v):
    return v.strip().lower() in MISSING_TOKENS


def participant_id_index(header, privacy_cols=()):
    """Column that identifies the participant: a strong name (participant,
    subject, ResponseId ...) first, a generic one (..._id) only if there is none;
    never a trial, item, node or other non-person ID, nor a personal-data column."""
    for rx in (PARTICIPANT_ID, PARTICIPANT_ID_WEAK):
        for j, h in enumerate(header):
            hs = h.strip()
            if hs and rx.search(hs) and not NOT_PARTICIPANT_ID.search(hs) and h not in privacy_cols:
                return j
    return None


def participant_values(header, body, cols, privacy_cols):
    """Group participant-level columns by participant ID.

    Returns (id_column, {id: {column: set(values)}}, list of within-participant problems).
    A column counts as participant-level when its name says so (age, sex, group, ...)
    or when it has more than two values and is constant for at least 80% of the
    participants with two or more rows: then the remaining ones are likely errors.
    """
    idx = participant_id_index(header, privacy_cols)
    if idx is None:
        return None, {}, []
    groups = defaultdict(lambda: defaultdict(set))
    for r in body:
        pid = r[idx].strip() if idx < len(r) else ""
        if not pid or is_missing(pid):
            continue
        for j, h in enumerate(header):
            if j != idx and j < len(r) and not is_missing(r[j]):
                groups[pid][h].add(r[j].strip())
    problems, person_cols = [], []
    n_rows = Counter(r[idx].strip() for r in body if idx < len(r))
    multi = [pid for pid in groups if n_rows[pid] > 1]
    for j, h in enumerate(header):
        if j == idx or h in privacy_cols:
            continue
        col = cols[j]
        varying = [pid for pid in multi if len(groups[pid].get(h, ())) > 1]
        named = bool(PERSON_LEVEL.match(canon_name(h)))
        generic = (col.get("n_unique", 0) > 2 and len(multi) >= 5 and 0 < len(varying) <= 0.2 * len(multi))
        if named or generic or not varying and multi:
            person_cols.append(h)
        if varying and (named or generic):
            ex = "; ".join(f"{pid}: {' / '.join(sorted(groups[pid][h]))}" for pid in varying[:5])
            problems.append(f"'{h}' should have one value per participant but has more for "
                            f"{len(varying)} of {len(multi)} participants ({ex})")
    by_id = {pid: {h: v for h, v in g.items() if h in person_cols} for pid, g in groups.items()}
    return header[idx], by_id, problems


def detect_decimal_comma(body, delimiter):
    """True when numbers are written 1,5 (possible only if commas do not separate fields)."""
    if delimiter in (None, "comma"):
        return False
    cells = [c.strip() for r in body[:500] for c in r]
    comma = sum(1 for c in cells if re.match(r"^-?\d+,\d+$", c))
    point = sum(1 for c in cells if re.match(r"^-?\d+\.\d+$", c))
    return comma > 0 and comma >= point


def scan_table(path, root, rows=None, info=None, sheet=None):
    """Check one table. `rows`/`info` come from load_any(); without them the
    file is read as delimited text."""
    name = rel(path, root) + (f" [{sheet}]" if sheet else "")
    if rows is None:
        try:
            rows, info = read_table(path)
        except Exception as e:  # noqa: BLE001
            return {"file": name, "error": f"could not read: {e}"}
    res = {"file": name, **{k: v for k, v in (info or {}).items() if not k.startswith("_")}}
    header, body, labels, layout = split_export(rows or [])
    if not header:
        res["error"] = "empty file"
        return res
    prog, evidence = detect_software(header)
    if not prog and labels:
        prog, evidence = detect_software(labels[0])
    if prog:
        res["software"] = prog
        res["software_evidence"] = evidence
    if layout or res.get("preamble"):
        if res.get("preamble"):
            layout["preamble_lines"] = res.pop("preamble")
        res["export_structure"] = layout
    res["n_rows"], res["n_cols"] = len(body), len(header)
    issues = []
    ragged = [i + 2 + len(labels) for i, r in enumerate(body) if len(r) != len(header)]
    if ragged:
        issues.append(f"{len(ragged)} rows with a different number of fields than the header (first lines: {ragged[:5]})")
    hdr_issues = []
    for i, h in enumerate(header):
        if not h.strip():
            hdr_issues.append(f"column {i + 1} has no name")
        else:
            p = [x for x in name_problems(h.strip(), is_dir=True) if not x.startswith("manual") and not x.startswith("looks like")]
            if p:
                hdr_issues.append(f"'{h}': {', '.join(p)}")
            if re.match(r"^\d", h.strip()):
                hdr_issues.append(f"'{h}': starts with a digit")
    dups = [h for h, c in Counter(header).items() if c > 1 and h.strip()]
    if dups:
        hdr_issues.append("duplicated column names: " + ", ".join(dups))
    if all(to_number(h) is not None for h in header if h.strip()):
        hdr_issues.append("the first row looks like data, not column names (no header?)")
    if hdr_issues:
        res["header_issues"] = hdr_issues
    tuples = [tuple(r) for r in body]
    dupe_rows = sum(c - 1 for c in Counter(tuples).values() if c > 1)
    if dupe_rows:
        issues.append(f"{dupe_rows} exact duplicate rows")
    decimal_comma = detect_decimal_comma(body, (info or {}).get("delimiter"))
    if decimal_comma:
        res["decimal_comma"] = True
    cols = []
    privacy = []
    for j, h in enumerate(header):
        vals = [r[j] if j < len(r) else "" for r in body]
        c = infer_column(h, vals, decimal_comma)
        pf = personal_flags(h, vals, labels[0][j] if labels and j < len(labels[0]) else "")
        if pf:
            privacy.append({"column": h, "flags": pf})
            c.pop("values", None)  # do not echo possibly identifying values
            c.pop("min", None)
            c.pop("max", None)
        cols.append(c)
    res["columns"] = cols
    exp_cols = [j for j, h in enumerate(header) if re.sub(r"[^a-z]", "", h.lower()) in ("experimentname", "experiment", "expname")]
    if exp_cols:
        res["_experiment"] = {r[j].strip().lower() for j in exp_cols for r in body if j < len(r) and r[j].strip()}
    if privacy:
        res["possible_personal_data"] = privacy
    id_col, by_id, within = participant_values(header, body, cols, {p["column"] for p in privacy})
    if id_col:
        res["participant_id_column"] = id_col
        res["n_participants"] = len(by_id)
        res["_by_id"] = by_id  # used for the cross-file comparison, removed before output
        word_ids = [k for k in by_id if not re.search(r"\d", k) and re.match(r"^[^\W\d_]{2,}([ _.\-][^\W\d_]+)*$", k)]
        if word_ids:
            privacy.append({"column": id_col, "flags": [
                f"{len(word_ids)} of {len(by_id)} IDs are letters only (initials or names?): check they cannot identify anyone"]})
            res["possible_personal_data"] = privacy
    else:
        personal_id = participant_id_index(header)
        if personal_id is not None and header[personal_id] in {p["column"] for p in privacy}:
            res["participant_id_column_is_personal"] = (
                f"'{header[personal_id]}' identifies the participants but holds personal data "
                "(e.g. a Prolific ID): replace it with a study code before sharing")
    if within:
        res["within_participant_inconsistencies"] = within
    res["issues"] = issues
    return res


# constant within a participant in one file, but different by design between
# files (session, task, timestamps, software settings): not compared across files
NOT_PERSON_ATTR = re.compile(r"(?i)(date|time|stamp|session|seed|version|task|node|schedule|experiment|exp_?name|file|path|"
                             r"index|count|status|key|code|frame|rate|zone|screen|spreadsheet|procedure|running|clock|"
                             r"browser|device|^os$|_os$|monitor|viewport|checkpoint|page|block|trial|run)")


def compare_participants(tables):
    """For each pair of files with a participant ID: which IDs are only in one of
    them, and whether shared participant-level columns (age, sex, ...) agree.
    Files from different experiments (different ExperimentName / expName) are
    not compared."""
    out = []
    with_ids = [t for t in tables if t.get("_by_id")]
    for i, a in enumerate(with_ids):
        for b in with_ids[i + 1:]:
            if a.get("_experiment") and b.get("_experiment") and not a["_experiment"] & b["_experiment"]:
                continue
            ida, idb = set(a["_by_id"]), set(b["_by_id"])
            key = {"a": lambda v: v, "b": lambda v: v}
            note = None
            if not ida & idb:
                # same participants written differently? ('sub-01' vs '1')
                def digits(v):
                    d = re.sub(r"\D", "", v).lstrip("0")
                    return d or v
                if {digits(v) for v in ida} & {digits(v) for v in idb}:
                    key = {"a": digits, "b": digits}
                    note = "IDs match only after removing prefixes and leading zeros: write them the same way in every file"
                else:
                    continue  # no participants in common: probably unrelated tables
            ma = {key["a"](k): k for k in ida}
            mb = {key["b"](k): k for k in idb}
            entry = {"files": [a["file"], b["file"]],
                     "id_columns": [a["participant_id_column"], b["participant_id_column"]],
                     "only_in_first": sorted(ma[k] for k in set(ma) - set(mb))[:15],
                     "only_in_second": sorted(mb[k] for k in set(mb) - set(ma))[:15]}
            if a["participant_id_column"] != b["participant_id_column"]:
                entry["note_names"] = "the same ID column has different names in the two files"
            if note:
                entry["note_format"] = note
            # shared participant-level columns
            cols_a = {canon_name(h): h for g in a["_by_id"].values() for h in g}
            cols_b = {canon_name(h): h for g in b["_by_id"].values() for h in g}
            disagreements = []
            for cn in sorted(set(cols_a) & set(cols_b)):
                ha, hb = cols_a[cn], cols_b[cn]
                if NOT_PERSON_ATTR.search(ha) and not PERSON_LEVEL.match(cn):
                    continue
                bad = []
                for k in sorted(set(ma) & set(mb)):
                    va = {v.lower() for v in a["_by_id"][ma[k]].get(ha, set())}
                    vb = {v.lower() for v in b["_by_id"][mb[k]].get(hb, set())}
                    na = {to_number(v) for v in va} if all(to_number(v) is not None for v in va) else va
                    nb = {to_number(v) for v in vb} if all(to_number(v) is not None for v in vb) else vb
                    if va and vb and na != nb:
                        bad.append(f"{ma[k]}: {' / '.join(sorted(va))} vs {' / '.join(sorted(vb))}")
                if bad:
                    disagreements.append({"columns": [ha, hb], "n_participants": len(bad), "examples": bad[:6]})
            if disagreements:
                entry["values_disagree"] = disagreements
            if entry["only_in_first"] or entry["only_in_second"] or disagreements or note or "note_names" in entry:
                out.append(entry)
    return out


def group_file_sets(tables):
    """Files in the same folder with (nearly) the same columns, e.g. one PsychoPy,
    OpenSesame or E-Prime file per participant, are one dataset: check them
    together, report them once. Files whose columns overlap by at least 80% go
    in the same set (a participant who never pressed a key may lack a column).

    Returns (tables to use for the cross-file and dictionary checks, entries for
    the report)."""
    clusters = []  # [parent, software, header set, [tables], extension]
    for t in tables:
        if "columns" not in t:
            continue
        parent = str(Path(t["file"].split(" [")[0]).parent)
        cols = {c["name"] for c in t["columns"]}
        ext = Path(t["file"].split(" [")[0]).suffix.lower()
        for cl in clusters:
            if cl[0] == parent and cl[1] == t.get("software") and cl[4] == ext and \
                    len(cols & cl[2]) >= 0.8 * len(cols | cl[2]):
                cl[3].append(t)
                cl[2] |= cols
                break
        else:
            clusters.append([parent, t.get("software"), set(cols), [t], ext])
    for_checks, for_report = [], []
    for parent, software, _, ts, _ in clusters:
        if len(ts) < 3:
            for_checks += ts
            for_report += ts
            continue
        exts = sorted({Path(t["file"].split(" [")[0]).suffix for t in ts})
        name = f"{parent}/*{exts[0] if len(exts) == 1 else ''} ({len(ts)} files with the same columns)"
        merged = defaultdict(dict)
        for t in ts:
            for pid, cols in t.get("_by_id", {}).items():
                for h, v in cols.items():
                    merged[pid].setdefault(h, set()).update(v)
        order, coldef, present_in = [], {}, Counter()
        for t in ts:
            for c in t["columns"]:
                present_in[c["name"]] += 1
                if c["name"] not in coldef:
                    order.append(c["name"])
                    coldef[c["name"]] = c
        exps = set().union(*[t.get("_experiment", set()) for t in ts])
        rep = {"file": name, "columns": [coldef[h] for h in order], "software": software, "_experiment": exps,
               "participant_id_column": ts[0].get("participant_id_column"), "_by_id": dict(merged)}
        for_checks.append(rep)
        per_file = {}
        for t in ts:
            found = t.get("header_issues", []) + t.get("issues", []) + t.get("within_participant_inconsistencies", []) + \
                [f"{c['name']}: {i}" for c in t["columns"] for i in c.get("issues", []) if not i.startswith("constant column")]
            if found:
                per_file[Path(t["file"]).name] = found
        types = defaultdict(set)
        for t in ts:
            for c in t["columns"]:
                types[c["name"]].add(c.get("type"))
        entry = {
            "file_set": name, "n_files": len(ts), "n_rows_total": sum(t.get("n_rows", 0) for t in ts),
            "files": sorted(t["file"] for t in ts),
            "software": software,
            "participant_id_column": rep["participant_id_column"], "n_participants": len(merged),
            "columns": [{"name": h, "types": sorted(x for x in types[h] if x)} for h in order],
            "issues_by_file": per_file,
            "possible_personal_data": sorted({p["column"] for t in ts for p in t.get("possible_personal_data", [])})}
        partial = {h: n for h, n in present_in.items() if n < len(ts)}
        if partial:
            entry["columns_not_in_every_file"] = {h: f"in {n} of {len(ts)} files" for h, n in partial.items()}
        structures = {json.dumps(t.get("export_structure", {}), sort_keys=True) for t in ts}
        if structures != {"{}"}:
            entry["export_structure"] = [json.loads(x) for x in structures]
        encodings = sorted({t.get("encoding") for t in ts if t.get("encoding")})
        entry["encoding"] = encodings[0] if len(encodings) == 1 else encodings
        for k in ("format", "rows_are"):
            vals = {str(t.get(k)) for t in ts if t.get(k)}
            if len(vals) == 1:
                entry[k] = vals.pop()
        for_report.append(entry)
    for_checks += [t for t in tables if "columns" not in t]
    for_report += [t for t in tables if "columns" not in t]
    return for_checks, for_report  # "_by_id" is removed in main() after the cross-file check


# --------------------------------------------------------------- dictionaries

def norm(s):
    return re.sub(r"[^a-z0-9]", "_", s.strip().lower()).strip("_")


def parse_dictionary(path, root):
    rows, info = read_table(path)
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return {"file": rel(path, root), "error": "empty"}
    header = [norm(h) for h in rows[0]]
    mapping = {}
    for std, aliases in DICT_ALIASES.items():
        for a in aliases:
            if norm(a) in header:
                mapping[std] = header.index(norm(a))
                break
    entries = {}
    if "variable_name" in mapping:
        for r in rows[1:]:
            k = mapping["variable_name"]
            if k < len(r) and r[k].strip():
                entries[r[k].strip()] = {std: (r[i].strip() if i < len(r) else "") for std, i in mapping.items()}
    return {"file": rel(path, root), "columns_present": [s for s in DICT_COLUMNS if s in mapping],
            "columns_missing": [s for s in DICT_COLUMNS if s not in mapping], "entries": entries}


def parse_codes(coding):
    """'1 = Psychology; 2 = Economics' -> {'1','2'}"""
    codes = set()
    for part in re.split(r"[;|\n]", coding):
        m = re.match(r"\s*([^=:]+?)\s*[=:]", part)
        if m:
            codes.add(m.group(1).strip())
    return codes


def compare_dictionary(table, dic):
    out = {"dictionary": dic["file"]}
    data_cols = [c["name"] for c in table.get("columns", [])]
    entries = dic.get("entries", {})
    out["data_columns_not_in_dictionary"] = [c for c in data_cols if c not in entries]
    out["dictionary_entries_not_in_data"] = [e for e in entries if e not in data_cols]
    near = []
    for c in out["data_columns_not_in_dictionary"]:
        for e in out["dictionary_entries_not_in_data"]:
            if norm(c) == norm(e):
                near.append(f"data '{c}' vs dictionary '{e}' (differ only in case/spacing)")
    if near:
        out["name_mismatches"] = near
    problems = []
    for col in table.get("columns", []):
        e = entries.get(col["name"])
        if not e:
            continue
        if not e.get("label"):
            problems.append(f"{col['name']}: no label")
        declared = (e.get("type") or "").lower()
        if declared not in ("string", "date", "numeric") and (
                col.get("type") in ("categorical", "boolean") or (col.get("type") == "integer" and col.get("values") and len(col["values"]) <= 6)):
            if "coding" in e and not e["coding"] and not e.get("allowed_values"):
                problems.append(f"{col['name']}: categorical-looking column without coding")
            codes = parse_codes(e.get("coding", ""))
            rng = re.match(r"^\s*(-?[\d.]+)\s*(?:-|–|to)\s*(-?[\d.]+)\s*$", e.get("allowed_values", ""))
            if rng and "min" in col and float(rng.group(1)) <= col["min"] and col["max"] <= float(rng.group(2)):
                codes = set()  # a scale: the coding labels its anchors, the range covers the values
            if codes and col.get("values"):
                extra = [v for v in col["values"] if v not in codes]
                if extra:
                    problems.append(f"{col['name']}: values in data not described in coding: {extra[:8]}")
        av = e.get("allowed_values", "")
        m = re.match(r"^\s*(-?[\d.]+)\s*(?:-|–|to|a|:)\s*(-?[\d.]+)\s*$", av)
        if m and "min" in col:
            lo, hi = float(m.group(1)), float(m.group(2))
            declared_missing = {to_number(x) for x in re.split(r"[;,|]", e.get("missing_values", "")) if to_number(x) is not None}
            if (col["min"] < lo and col["min"] not in declared_missing) or (col["max"] > hi and col["max"] not in declared_missing):
                problems.append(f"{col['name']}: data range {col['min']}–{col['max']} outside allowed_values {av}")
        if col.get("n_missing") and "missing_values" in e and not e["missing_values"]:
            problems.append(f"{col['name']}: has {col['n_missing']} missing values but missing_values is empty")
    out["entry_problems"] = problems
    return out


# ----------------------------------------------------------------------- code

def scan_code(path, root, all_files_rel):
    text = path.read_text(errors="replace")
    ext = path.suffix.lower()
    lang = CODE_EXT.get(ext, "?")
    res = {"file": rel(path, root), "language": lang, "lines": text.count("\n") + 1}
    if ext == ".ipynb":
        try:
            nb = json.loads(text)
            text = "\n".join("".join(c.get("source", [])) for c in nb.get("cells", []) if c.get("cell_type") == "code")
            kern = nb.get("metadata", {}).get("kernelspec", {}).get("language", "python")
            lang = "R" if kern.lower() == "r" else "Python"
            res["language"] = f"Jupyter ({lang})"
        except Exception:  # noqa: BLE001
            pass
    code = "\n".join(l for l in text.splitlines() if not l.lstrip().startswith("#") or ext in (".py", ".ipynb") and l.lstrip().startswith("#!"))
    pkgs = set()
    if lang in ("R", "R Markdown", "Quarto") or res["language"] == "Jupyter (R)":
        for m in re.finditer(r"\b(?:library|require|requireNamespace)\s*\(\s*[\"']?([A-Za-z][A-Za-z0-9.]*)", code):
            pkgs.add(m.group(1))
        for m in re.finditer(r"\b([A-Za-z][A-Za-z0-9.]*)::", code):
            pkgs.add(m.group(1))
        for m in re.finditer(r"p_load\s*\(([^)]*)\)", code):
            for p in re.split(r"[,\s]+", m.group(1)):
                p = p.strip("\"' ")
                if re.match(r"^[A-Za-z][A-Za-z0-9.]*$", p) and "=" not in p:
                    pkgs.add(p)
        pkgs -= R_BASE
        if ext == ".rmd" or ext == ".qmd":
            pkgs.add("rmarkdown" if ext == ".rmd" else "knitr")
    elif lang in ("Python",) or res["language"] == "Jupyter (Python)":
        std = set(getattr(sys, "stdlib_module_names", ()))
        local = {Path(f).stem for f in all_files_rel if f.endswith(".py")}
        for m in re.finditer(r"^\s*(?:import\s+([\w.]+(?:\s*,\s*[\w.]+)*)|from\s+([\w.]+)\s+import)", code, re.M):
            names = (m.group(1) or m.group(2)).split(",")
            for nme in names:
                top = nme.strip().split(".")[0].split(" as ")[0].strip()
                if top and top not in std and top not in local:
                    pkgs.add(top)
    res["packages"] = sorted(pkgs)
    issues = []
    for m in re.finditer(r"setwd\s*\(([^)]*)\)", code):
        issues.append(f"setwd({m.group(1).strip()}) — ties the script to one computer")
    abs_paths = set(re.findall(r"[\"']((?:[A-Za-z]:[\\/]|/Users/|/home/|~/|\\\\[A-Za-z])[^\"']*)[\"']", code))
    for p in sorted(abs_paths)[:10]:
        issues.append(f"absolute path '{p}'")
    if re.search(r"rm\s*\(\s*list\s*=\s*ls\s*\(", code):
        issues.append("rm(list = ls()) — does not give a clean session; restart R instead")
    if re.search(r"install\.packages\s*\(|pip install|!pip", code):
        issues.append("installs packages inside the script")
    # files the script writes (so a later script reading them is not a "missing file")
    res["writes"] = sorted({Path(m.group(1)).name for m in re.finditer(
        r"(?:write[._]\w*|to_csv|to_excel|saveRDS|save|ggsave|fwrite|sink|savefig|png|pdf|jpeg)\s*\([^)]*?[\"']([^\"'\n]+\.\w+)[\"']", code)})
    # file references that do not exist
    missing = []
    for m in re.finditer(r"[\"']([^\"'\n]+\.(?:csv|tsv|txt|xlsx|xls|sav|rds|RData|rda|json|dta|mat|feather|parquet))[\"']", code):
        ref = m.group(1)
        if re.match(r"^(?:[A-Za-z]:[\\/]|/|~|https?:)", ref) or "*" in ref or "%" in ref or "{" in ref:
            continue
        cand = [root / ref, path.parent / ref, root / ref.lstrip("./")]
        # a bare file name may be one piece of here("data", "x.csv") or file.path(...)
        if "/" not in ref and any(Path(f).name == ref for f in all_files_rel):
            continue
        if not any(c.exists() for c in cand):
            missing.append(ref)
    if missing:
        res["referenced_files_not_found"] = sorted(set(missing))[:20]
    res["issues"] = issues
    res["uses_here"] = bool(re.search(r"\bhere\s*::\s*here|library\s*\(\s*here\s*\)", code))
    if lang in ("R", "R Markdown", "Quarto") or res["language"] == "Jupyter (R)":
        rnd, seed = RANDOM_R, SEED_R
    elif lang == "MATLAB":
        rnd, seed = RANDOM_M, SEED_M
    else:
        rnd, seed = RANDOM_PY, SEED_PY
    funcs = sorted({re.sub(r"[\s(]+$", "", m.group(0)).lstrip("$") for m in rnd.finditer(code)})
    seeds = sorted({m.group(0).strip()[:40] for m in seed.finditer(code)})
    res["sets_seed"] = bool(seeds)
    res["random_calls"] = bool(funcs)
    if funcs:
        res["randomness"] = {"functions": funcs[:12], "seed": seeds[:5] or None}
        if not seeds:
            issues.append("uses random numbers (" + ", ".join(funcs[:5]) + ") without fixing a seed: "
                          "the results change slightly at every run")
    return res


def parse_env(root):
    env = {}
    p = root / "renv.lock"
    if p.exists():
        try:
            lock = json.loads(p.read_text())
            env["renv.lock"] = {"R": lock.get("R", {}).get("Version"), "packages": sorted(lock.get("Packages", {}))}
        except Exception as e:  # noqa: BLE001
            env["renv.lock"] = {"error": str(e)}
    p = root / "requirements.txt"
    if p.exists():
        env["requirements.txt"] = [re.split(r"[<>=!~\[; ]", l.strip())[0] for l in p.read_text().splitlines() if l.strip() and not l.startswith("#")]
    for name in ("environment.yml", "environment.yaml", "DESCRIPTION", "pyproject.toml", "sessionInfo.txt", "session_info.txt"):
        if (root / name).exists():
            env[name] = "present"
    return env


# -------------------------------------------------------- manuscript, psych-ds

MANUSCRIPT_NAME = re.compile(r"(?i)(manuscript|manoscritto|paper|preprint|article|articolo|draft|submission|thesis|tesi)")
MANUSCRIPT_EXT = {".pdf", ".docx", ".doc", ".odt", ".md", ".qmd", ".rmd", ".tex", ".html"}
PSYCHDS_NAME = re.compile(r"^([a-z]+-[a-zA-Z0-9]+_)+data\.csv$")
PSYCHDS_KEYWORDS = {"study", "site", "subject", "session", "task", "condition", "trial", "stimulus", "description"}


def find_manuscripts(root, files):
    out = []
    for f in files:
        r = rel(f, root)
        if f.suffix.lower() not in MANUSCRIPT_EXT or re.match(r"(?i)^(readme|leggimi|licen[cs]e|changelog)", f.name):
            continue
        in_docs = any(p.lower() in {"docs", "doc", "manuscript", "manoscritto", "paper", "preprint", "article", "articolo"}
                      for p in Path(r).parts[:-1])
        if in_docs or MANUSCRIPT_NAME.search(f.stem):
            out.append({"file": r, "format": f.suffix.lower(), "size_kb": round(f.stat().st_size / 1024, 1)})
    return out


def psychds_check(root, tables):
    """The minimum rules of Psych-DS, as listed on the booklet's page."""
    res = {"data_dir": (root / "data").is_dir(), "problems": []}
    dd = root / "dataset_description.json"
    if not dd.exists():
        res["problems"].append("no dataset_description.json in the root")
        meta = None
    else:
        try:
            meta = json.loads(dd.read_text())
        except json.JSONDecodeError as e:
            res["problems"].append(f"dataset_description.json is not valid JSON: {e}")
            meta = None
    if (root / "data" / "dataset_description.json").exists():
        res["problems"].append("dataset_description.json is inside data/: it must be in the root")
    in_data = [t for t in tables if t["file"].startswith("data/") and not t["file"].startswith("data/raw/")]
    bad_names = [t["file"] for t in in_data if not PSYCHDS_NAME.match(Path(t["file"]).name)]
    if bad_names:
        res["problems"].append("files in data/ not named key-value_..._data.csv: " + ", ".join(bad_names))
    unknown = sorted({k for t in in_data if PSYCHDS_NAME.match(Path(t["file"]).name)
                      for k in re.findall(r"([a-z]+)-[a-zA-Z0-9]+_", Path(t["file"]).name) if k not in PSYCHDS_KEYWORDS})
    if unknown:
        res["warnings"] = ["keys not in the official Psych-DS list (" + ", ".join(sorted(PSYCHDS_KEYWORDS)) + "): " + ", ".join(unknown)]
    if not res["data_dir"]:
        res["problems"].append("no data/ folder in the root")
    if meta is not None:
        for req in ("name", "description", "variableMeasured"):
            if not meta.get(req):
                res["problems"].append(f"dataset_description.json: required field '{req}' missing or empty")
        declared = {v if isinstance(v, str) else v.get("name") for v in meta.get("variableMeasured", [])}
        cols = {c["name"] for t in in_data if PSYCHDS_NAME.match(Path(t["file"]).name) for c in t.get("columns", [])}
        if cols - declared:
            res["problems"].append("columns not in variableMeasured: " + ", ".join(sorted(cols - declared)))
        if declared - cols and cols:
            res["problems"].append("variableMeasured entries that are not columns of any data file: " + ", ".join(sorted(declared - cols)))
        res["n_variables_described"] = sum(1 for v in meta.get("variableMeasured", []) if isinstance(v, dict) and v.get("description"))
    res["valid_minimum"] = not res["problems"]
    return res


# --------------------------------------------------------------------- readme

def find_readme(root):
    for f in sorted(root.iterdir()):
        if f.is_file() and re.match(r"(?i)^(readme|leggimi)(\.(md|txt|rst|markdown|html|pdf|docx))?$", f.name):
            return f
    return None


def scan_readme(path, root, files_rel):
    res = {"file": rel(path, root), "format": path.suffix.lower() or "(none)"}
    if path.suffix.lower() in (".pdf", ".docx", ".html"):
        res["issues"] = [f"README in {path.suffix} — use Markdown or plain text"]
        return res
    text = path.read_text(errors="replace")
    res["n_words"] = len(text.split())
    lines = [l for l in text.splitlines() if l.strip()]
    res["first_line"] = lines[0][:120] if lines else ""
    headings = [l.strip("# ").strip() for l in text.splitlines() if re.match(r"^#{1,6}\s", l)]
    res["headings"] = headings[:40]
    found = {}
    for sec, pat in README_SECTIONS.items():
        if pat is None:
            found[sec] = bool(lines)
            continue
        in_head = any(re.search(pat, h) for h in headings)
        in_text = bool(re.search(pat, text))
        found[sec] = "heading" if in_head else ("mentioned" if in_text else False)
    res["sections"] = found
    res["orcid_ids"] = re.findall(r"\d{4}-\d{4}-\d{4}-\d{3}[\dX]", text)
    res["dois"] = sorted(set(re.findall(r"10\.\d{4,9}/[^\s)\]>\"']+", text)))[:10]
    res["placeholders"] = sorted(set(re.findall(r"(?i)(TODO|TBD|XXXX+|lorem ipsum|\[insert[^\]]*\])", text)))
    mentioned, not_mentioned = [], []
    for f in files_rel:
        if Path(f).name.lower().startswith(("readme", "leggimi")):
            continue
        if Path(f).name in text or f in text or Path(f).stem in text:
            mentioned.append(f)
        else:
            parent = str(Path(f).parent)
            if parent != "." and (parent + "/" in text or parent + "\\" in text or
                                  (Path(parent).name not in ("data", "raw", "scripts", "code") and Path(parent).name + "/" in text)):
                mentioned.append(f)
            else:
                not_mentioned.append(f)
    res["files_mentioned"] = len(mentioned)
    res["files_not_mentioned"] = not_mentioned[:60]
    return res


# -------------------------------------------------------------------- license

def detect_licenses(text):
    found = []
    for spdx, pats in LICENSE_SIGNATURES:
        for p in pats:
            if re.search(p, text):
                if spdx == "GPL-3.0" and ("AGPL-3.0" in found or "LGPL-3.0" in found) and not re.search(r"(?<![AL])GPL", text):
                    break
                if spdx.startswith("CC-BY") and any(f.startswith("CC-BY") and f != spdx and len(f) > len(spdx) for f in found) and not re.search(pats[0], text):
                    break
                found.append(spdx)
                break
    return found


def scan_licenses(root, files):
    out = {"files": []}
    for f in files:
        if re.match(r"(?i)^(licen[cs]e|copying|copyright)", f.name):
            txt = f.read_text(errors="replace")[:200000]
            out["files"].append({"file": rel(f, root), "detected": detect_licenses(txt),
                                 "n_chars": len(txt),
                                 "looks_like_full_text": len(txt) > 900})
    return out


# ----------------------------------------------------------------------- main

def scan(root):
    """Scan the folder `root` and return the report as a dict."""
    root = Path(root).expanduser().resolve()
    files, dirs = walk(root)
    files_rel = [rel(f, root) for f in files]
    report = {"project": str(root), "n_files": len(files), "n_dirs": len(dirs)}

    # inventory
    inv = []
    total = 0
    for f in files:
        size = f.stat().st_size
        total += size
        inv.append({"path": rel(f, root), "size_kb": round(size / 1024, 1)})
    report["total_size_mb"] = round(total / 1e6, 2)
    report["inventory"] = inv[:400]
    if len(inv) > 400:
        report["inventory_truncated"] = len(inv)
    report["root_entries"] = sorted(p.name + ("/" if p.is_dir() else "") for p in root.iterdir()
                                    if not p.name.startswith(".") and p.name not in SKIP_FILES)

    # naming
    naming = []
    for d in dirs:
        p = name_problems(Path(d).name, is_dir=True)
        if p:
            naming.append({"path": d + "/", "problems": p})
    for f in files:
        p = name_problems(f.name)
        if p:
            naming.append({"path": rel(f, root), "problems": p})
    report["naming"] = {"problems": naming, "numbering": check_numbering(files, root)}

    # formats
    fmt = {"proprietary_data": [], "proprietary_docs": [], "open_tabular": [], "other_binary": []}
    stems_open = {(f.parent, f.stem.lower()) for f in files if f.suffix.lower() in OPEN_TABULAR}
    stems_open_any = {f.stem.lower() for f in files if f.suffix.lower() in OPEN_TABULAR}
    for f in files:
        ext = f.suffix.lower()
        r = rel(f, root)
        if ext in PROPRIETARY_DATA:
            fmt["proprietary_data"].append({"file": r, "format": PROPRIETARY_DATA[ext],
                                            "open_copy_exists": (f.parent, f.stem.lower()) in stems_open or f.stem.lower() in stems_open_any,
                                            "in_raw_folder": "raw" in [p.lower() for p in Path(r).parts[:-1]]})
        elif ext in PROPRIETARY_DOCS:
            fmt["proprietary_docs"].append({"file": r, "format": PROPRIETARY_DOCS[ext]})
        elif ext in OPEN_TABULAR:
            fmt["open_tabular"].append(r)
    report["formats"] = fmt

    # structure
    lower_dirs = {d.lower() for d in dirs}
    top_dirs = {d.lower() for d in dirs if "/" not in d and os.sep not in d}
    report["structure"] = {
        "readme": bool(find_readme(root)),
        "license_file": any(re.match(r"(?i)^(licen[cs]e|copying)", f.name) for f in files if f.parent == root),
        "citation_cff": (root / "CITATION.cff").exists(),
        "rproj": [rel(f, root) for f in files if f.suffix.lower() == ".rproj"],
        "dataset_description_json": (root / "dataset_description.json").exists(),
        "data_dir": sorted(d for d in top_dirs if d in {"data", "dati", "dataset", "datasets"}),
        "raw_dir": sorted(d for d in lower_dirs if Path(d).name in {"raw", "raw_data", "rawdata", "grezzi", "sourcedata", "original"}),
        "code_dir": sorted(d for d in top_dirs if d in {"scripts", "script", "code", "codice", "analysis", "analyses", "analisi", "r", "src", "notebooks"}),
        "output_dir": sorted(d for d in top_dirs if d in {"output", "outputs", "results", "risultati", "figures", "figure", "plots", "tables"}),
        "materials_dir": sorted(d for d in top_dirs if d in {"materials", "materiali", "stimuli", "stimoli", "questionnaires", "questionari", "task", "experiment", "esperimento"}),
        "docs_dir": sorted(d for d in top_dirs if d in {"docs", "doc", "documents", "documenti", "manuscript", "paper", "preprint"}),
        "files_in_root": [f.name for f in files if f.parent == root],
        "empty_dirs": [d + "/" for d in dirs if not any(r.startswith(d + "/") for r in files_rel)],
        "max_depth": max((len(Path(r).parts) for r in files_rel), default=0),
        "environment_files": [e for e in ENV_FILES if (root / e).exists()],
    }

    # readme
    rd = find_readme(root)
    report["readme"] = scan_readme(rd, root, files_rel) if rd else None
    nested = [r for r in files_rel if re.match(r"(?i)^(readme|leggimi)", Path(r).name) and "/" in r]
    if nested:
        report["nested_readmes"] = nested

    # licenses
    lic = scan_licenses(root, files)
    if rd and rd.suffix.lower() in (".md", ".txt", "", ".rst", ".markdown"):
        rtxt = rd.read_text(errors="replace")
        m = re.search(r"(?is)#+\s*(licen[sc]e|licenza)[^\n]*\n(.*?)(?=\n#{1,6}\s|\Z)", rtxt)
        lic["readme_license_section"] = m.group(2).strip()[:1500] if m else None
        lic["readme_detected"] = detect_licenses(m.group(2)) if m else detect_licenses(rtxt)
    for meta in ("CITATION.cff", "DESCRIPTION", "dataset_description.json", "codemeta.json", ".zenodo.json"):
        p = root / meta
        if p.exists():
            mm = re.search(r"(?im)^\s*\"?licen[sc]e\"?\s*[:=]\s*(.+)$", p.read_text(errors="replace"))
            if mm:
                lic.setdefault("metadata_files", {})[meta] = mm.group(1).strip().strip('",')
    report["licenses"] = lic

    # data + dictionaries
    tables, dicts, not_read, open_stems = [], {}, [], set()
    for f in files:
        suf, r = f.suffix.lower(), rel(f, root)
        if suf in OPEN_TABULAR and DICT_NAME.search(f.stem):
            try:
                d = parse_dictionary(f, root)
            except Exception as e:  # noqa: BLE001
                d = {"file": r, "error": str(e)}
            if "error" in d or "variable_name" in d.get("columns_present", []):
                dicts[r] = d
                continue
            # named like a dictionary, but no column of variable names: data
        if not (suf in OPEN_TABULAR or suf in TEXT_DATA or suf in JSON_DATA or suf in PROPRIETARY_DATA):
            continue
        try:
            found = load_any(f)
        except Exception as e:  # noqa: BLE001
            found = [(None, None, {"not_read": f"could not read: {e}"})]
        for sheet, rows, info in found:
            if rows is None:
                not_read.append({"file": r, "reason": info.get("not_read")})
                continue
            tables.append(scan_table(f, root, rows, info, sheet))
            if suf not in PROPRIETARY_DATA:
                open_stems.add(f.stem.lower())
    if not_read:
        report["data_not_read"] = not_read
    # proprietary files: is there an open copy, and does it hold the same table?
    for e in report["formats"]["proprietary_data"]:
        stem = Path(e["file"]).stem.lower()
        if stem in open_stems:
            e["open_copy_exists"] = True
        mine = [t for t in tables if t["file"].split(" [")[0] == e["file"] and "columns" in t]
        copies = [t for t in tables if "columns" in t and Path(t["file"].split(" [")[0]).stem.lower() == stem
                  and Path(t["file"].split(" [")[0]).suffix.lower() not in PROPRIETARY_DATA]
        if mine and copies:
            a, b = mine[0], copies[0]
            same_cols = [c["name"] for c in a["columns"]] == [c["name"] for c in b["columns"]]
            e["open_copy_check"] = {"copy": b["file"], "same_columns": same_cols,
                                    "rows": f"{a.get('n_rows')} here, {b.get('n_rows')} in the copy"}
    other_dicts = [r for r in files_rel if DICT_NAME.search(Path(r).stem) and Path(r).suffix.lower() not in OPEN_TABULAR]
    all_tables = tables
    tables, report["data"] = group_file_sets(tables)
    # shared columns with different types across files
    by_name = defaultdict(set)
    for t in tables:
        for c in t.get("columns", []):
            if c.get("type") not in (None, "empty"):
                by_name[c["name"]].add((c["type"], t["file"]))
    incons = {}
    for n, s in by_name.items():
        types = {t for t, _ in s}
        if len(types) > 1 and not types <= {"integer", "numeric"} and not types <= {"string", "categorical"}:
            incons[n] = sorted(f"{f}: {t}" for t, f in s)
    if incons:
        report["cross_file_type_inconsistencies"] = incons
    report["cross_file_participants"] = compare_participants(tables)
    for t in tables + all_tables + report["data"]:
        t.pop("_by_id", None)
        t.pop("_experiment", None)

    dict_report = {"dictionary_files": [{k: v for k, v in d.items() if k != "entries"} | {"n_entries": len(d.get("entries", {}))} for d in dicts.values()],
                   "non_csv_dictionaries": other_dicts, "coverage": []}
    psychds = root / "dataset_description.json"
    psychds_vars = []
    if psychds.exists():
        try:
            dd = json.loads(psychds.read_text())
            psychds_vars = [v if isinstance(v, str) else v.get("name") for v in dd.get("variableMeasured", [])]
            dict_report["dataset_description_variables"] = psychds_vars
        except Exception as e:  # noqa: BLE001
            dict_report["dataset_description_error"] = str(e)
    for t in tables:
        if "columns" not in t:
            continue
        tf = Path(t["file"])
        match = None
        tstem = re.sub(r"(?i)[_-]?data$", "", tf.stem).lower()
        stems = {dname: DICT_NAME.sub("", Path(dname).stem).strip("_- ").lower() for dname in dicts}
        # an exact name first (x_data.csv -> x_dictionary.csv), then a shared beginning
        for exact in (True, False):
            for dname, dstem in stems.items():
                if not dstem:
                    continue
                if (dstem in (tf.stem.lower(), tstem)) if exact else \
                        (tf.stem.lower().startswith(dstem) or dstem.startswith(tstem)):
                    match = dicts[dname]
                    break
            if match is not None:
                break
        if match is None:
            # a single generic dictionary that covers this file's columns
            cols = {c["name"] for c in t["columns"]}
            best = max(dicts.values(), key=lambda d: len(cols & set(d.get("entries", {}))), default=None)
            if best and len(cols & set(best.get("entries", {}))) >= max(1, len(cols) // 2):
                match = best
        if match is None and psychds_vars:
            cols = [c["name"] for c in t["columns"]]
            dict_report["coverage"].append({
                "data": t["file"], "dictionary": "dataset_description.json (variableMeasured: names only)",
                "data_columns_not_in_dictionary": [c for c in cols if c not in psychds_vars]})
        elif match is None:
            dict_report["coverage"].append({"data": t["file"], "dictionary": None})
        else:
            dict_report["coverage"].append({"data": t["file"], **compare_dictionary(t, match)})
    report["dictionaries"] = dict_report

    # experiment files (the instrument) and analysis code
    experiment = []
    for f in files:
        suf = f.suffix.lower()
        if suf in EXPERIMENT_EXT or (suf in (".py", ".js", ".html", ".htm", ".m") and f.stat().st_size < 5_000_000):
            role = experiment_role(f, "" if suf in EXPERIMENT_EXT else f.read_text(errors="replace"))
            if role:
                experiment.append({"file": rel(f, root), "what": role})
    report["experiment_files"] = experiment
    exp_set = {e["file"] for e in experiment}
    code_files = [f for f in files if f.suffix.lower() in CODE_EXT and rel(f, root) not in exp_set
                  and not (f.suffix.lower() == ".m" and f.stat().st_size == 0)]
    code = [scan_code(f, root, files_rel) for f in code_files]
    written = {w for c in code for w in c.get("writes", [])}
    for c in code:
        if "referenced_files_not_found" in c:
            c["referenced_files_not_found"] = [r for r in c["referenced_files_not_found"] if Path(r).name not in written]
            if not c["referenced_files_not_found"]:
                del c["referenced_files_not_found"]
    report["code"] = {"files": code, "environment": parse_env(root)}
    report["manuscripts"] = find_manuscripts(root, files)
    report["psychds"] = psychds_check(root, all_tables)
    allp = sorted({p for c in code for p in c.get("packages", [])})
    report["code"]["all_packages"] = allp
    lock = report["code"]["environment"].get("renv.lock", {})
    if isinstance(lock, dict) and lock.get("packages"):
        report["code"]["packages_not_in_renv_lock"] = [p for p in allp if p not in lock["packages"]]

    archives = [rel(f, root) for f in files if f.suffix.lower() in ARCHIVES]
    if archives:
        report["archives"] = {"files": archives, "note": "contents not checked: extract them to a temporary folder and scan that"}
    report["materials_present"] = {
        "data": bool(tables) or bool(fmt["proprietary_data"]) or bool(not_read),
        "code": bool(code),
        "manuscript": bool(report["manuscripts"]),
        "experiment": bool(experiment),
    }
    return report


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("folder")
    ap.add_argument("--out")
    args = ap.parse_args()
    root = Path(args.folder).expanduser().resolve()
    if not root.is_dir():
        sys.exit(f"Not a folder: {root}")
    report = scan(root)
    text = json.dumps(report, indent=1, ensure_ascii=False, default=str)
    if args.out:
        Path(args.out).write_text(text)
        print(f"Report written to {args.out} ({len(text)} chars)")
    else:
        print(text)


if __name__ == "__main__":
    main()
